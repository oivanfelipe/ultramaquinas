"""
Generates output/seo_metadata.xlsx with optimized SEO Titles and Meta Descriptions
for the 24 Verão campaign products.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PRODUCTS = [
    {
        "url": "https://www.ultramaquinas.com.br/produto/cadeira-dobravel-alvorada-para-camping-e-pesca-verde-290380-nautika-76491",
        "keyword": "cadeira de pesca dobrável",
        "seo_title": "Cadeira de Pesca Dobrável Alvorada Nautika | Ultra Máquinas",
        "meta_description": "Leve e resistente: armação reforçada em aço, bolso lateral e assento em brim reforçado. Para camping e pesca. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/mesa-dobravel-de-aluminio-azul-com-4-banquetas-para-camping-gh200-globalmix-75620",
        "keyword": "mesa dobravel camping",
        "seo_title": "Mesa Camping Alumínio com 4 Banquetas GH200 | Ultra Máquinas",
        "meta_description": "Mesa de camping com 4 banquetas em alumínio: montagem rápida, design compacto, perfeita para acampar ou praia. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/barraca-infantil-2m-x-1m-c-bolsa-e-bolinhas-bw065-importway-77247",
        "keyword": "barraquinha para criança brincar",
        "seo_title": "Barraca Infantil 2m com Bolinhas e Bolsa | Ultra Máquinas",
        "meta_description": "Barraca infantil 2x1m com bolinhas inclusas e bolsa para transporte — diversão garantida dentro de casa. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/ventilador-de-parede-premium-60cm-preto-bivolt-73-6425-venti-delta-78128",
        "keyword": "ventilador delta premium 60cm",
        "seo_title": "Ventilador Delta Premium 60cm Bivolt Preto | Ultra Máquinas",
        "meta_description": "Alto desempenho: ventilador de parede Venti Delta Premium 60cm, bivolt, 6 pás — ideal para ambientes grandes. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/esteira-de-praia-palha-natural-cores-sortidas-180x70cm-bel-76401",
        "keyword": "esteira de praia palha",
        "seo_title": "Esteira de Praia Palha Natural 180x70cm Bel | Ultra Máquinas",
        "meta_description": "Esteira de praia em palha natural 180x70cm — resistente, sustentável e fácil de enrolar. Cores sortidas. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/cadeira-de-praia-alta-em-aluminio-lazy-mormaii-preta-bel-77602",
        "keyword": "cadeira de praia alta alumínio",
        "seo_title": "Cadeira Praia Alta Lazy Mormaii Alumínio | Ultra Máquinas",
        "meta_description": "Cadeira de praia alta em alumínio resistente, design Lazy by Mormaii — estrutura leve com conforto superior. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/cama-de-solteiro-dobravel-kayman-camping-bege-291030-ntk-78023",
        "keyword": "cama de armar solteiro",
        "seo_title": "Cama Dobrável Solteiro Kayman Camping NTK | Ultra Máquinas",
        "meta_description": "Cama de camping dobrável para solteiro em estrutura de aço tubular resistente — fácil de montar e transportar. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/mesa-dobravel-com-ajuste-de-altura-robust-80cm-ntk-preto-291101-nautika-78828",
        "keyword": "mesa ntk robust 80cm",
        "seo_title": "Mesa NTK Robust 80cm com Ajuste de Altura | Ultra Máquinas",
        "meta_description": "Mesa dobrável NTK Robust com ajuste de altura até 80cm — pés reguláveis, tampo em MDF reforçado e estrutura em aço. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/cadeira-de-praia-copacabana-reclinavel-5-posicoes-rosa-cad0683-botafogo-76542",
        "keyword": "cadeira de praia botafogo copacabana",
        "seo_title": "Cadeira de Praia Reclinável Copacabana Rosa Botafogo",
        "meta_description": "Cadeira de praia rosa com 5 posições de reclinação — armação em alumínio leve e tecido resistente à areia. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/cadeira-praia-reclinavel-5-posicoes-botafogo-74360",
        "keyword": "cadeira de praia botafogo",
        "seo_title": "Cadeira de Praia Botafogo Reclinável | Ultra Máquinas",
        "meta_description": "Cadeira de praia Botafogo com 5 posições de reclinação — estrutura em alumínio leve e fácil de transportar. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/lavadora-de-alta-vazao-2-0cv-el-4000v2-a-39453-8-eletroplas-76252",
        "keyword": "lavadora eletroplas 2cv",
        "seo_title": "Lavadora Eletroplas 2CV Alta Pressão | Ultra Máquinas",
        "meta_description": "Potência 2CV para limpeza pesada: lavadora Eletroplas EL-4000V2 com alta vazão e motor elétrico robusto. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/mesa-dobravel-camping-nautika-robust-80cm-78121",
        "keyword": "mesa dobravel camping nautika",
        "seo_title": "Mesa Dobrável Camping Nautika Robust 80cm | Ultra Máquinas",
        "meta_description": "Mesa dobrável de camping Nautika Robust com 80cm de altura — estrutura em aço leve e fácil de montar. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/lavadora-de-alta-jaguar-turbo-pressao-2175-libras-127v-b8-065-0009-lavor-77105",
        "keyword": "lavadora jaguar turbo 2175",
        "seo_title": "Lavadora Jaguar Turbo 2175 Libras Lavor | Ultra Máquinas",
        "meta_description": "Lavadora Lavor Jaguar Turbo com 2175 libras de pressão e 127V — potência profissional para limpeza pesada. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/lavadora-de-alta-pressao-hd-498-1800w-14237050-karcher-75780",
        "keyword": "lavadora karcher hd 498",
        "seo_title": "Lavadora Karcher HD 498 1800W Alta Pressão | Ultra Máquinas",
        "meta_description": "Lavadora Kärcher HD 498 com 1800W de potência e alta pressão — desempenho profissional para superfícies externas. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/ombrellone-central-pisa-articulado-em-aluminio-poliester-2-50m-preto-14507-bel-fix-78575",
        "keyword": "ombrellone articulado 2,50m",
        "seo_title": "Ombrellone Pisa Articulado 2,50m Alumínio | Ultra Máquinas",
        "meta_description": "Ombrellone articulado com haste de alumínio 2,50m e tecido em poliéster preto — ideal para jardim e piscina. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/tenda-gazebo-dobravel-3m-x-3m-poliester-bege-bel-331500-bel-fix-77202",
        "keyword": "gazebo dobrável 3x3",
        "seo_title": "Gazebo Dobrável 3x3m Bege Bel Fix | Ultra Máquinas",
        "meta_description": "Gazebo dobrável 3x3m em poliéster bege resistente — montagem sem ferramentas e estrutura em aço tratado. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/caixa-termica-19-litros-acai-71909-bel-77810",
        "keyword": "caixa térmica 19 litros",
        "seo_title": "Caixa Térmica Açaí 19 Litros Bel 71909 | Ultra Máquinas",
        "meta_description": "Caixa térmica Açaí 19 litros — mantém sua bebida gelada por horas com isolamento térmico de alta performance. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/tenda-gazebo-tubular-2m-polietileno-azul-belfix-301302-74641",
        "keyword": "tenda gazebo 2m",
        "seo_title": "Gazebo Tubular 2m Polietileno Azul BelFix | Ultra Máquinas",
        "meta_description": "Tenda gazebo tubular 2m em polietileno azul — estrutura leve e montagem rápida para eventos ao ar livre. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/ombrellone-suspenso-giratorio-360%C2%B0-buzios-3m-marrom-c-base-50-lts-891012-bel-77818",
        "keyword": "ombrellone suspenso giratório 3m",
        "seo_title": "Ombrellone Suspenso Giratório 360° Búzios 3m + Base Bel",
        "meta_description": "Ombrellone suspenso giratório 360°, 3m de cobertura com base de água e areia 50L incluída — sombra total. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/guarda-sol-bagum-haste-de-aluminio-2-00m-azul-royal-10602-bel-74687",
        "keyword": "guarda-sol alumínio 2m",
        "seo_title": "Guarda-Sol Bagum Alumínio 2m Azul Royal Bel | Ultra Máquinas",
        "meta_description": "Guarda-sol Bagum com haste de alumínio 2m e tecido azul royal — leve e resistente para praia ou piscina. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/ombrellone-central-retangular-quiosque-oversize-4-50m-marrom-12412-bel-fix-78574",
        "keyword": "ombrellone retangular 4,50m",
        "seo_title": "Ombrellone Retangular 4,50m Oversize Marrom Bel Fix",
        "meta_description": "Ombrellone retangular de 4,50m tipo quiosque — ideal para áreas extensas com cobertura ampla e estrutura robusta. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/base-para-guarda-sol-e-ombrellone-agua-areia-18kg-preta-70027107-bel-fix-76391",
        "keyword": "base para guarda-sol 18kg",
        "seo_title": "Base Guarda-Sol 18kg Água/Areia Bel Fix | Ultra Máquinas",
        "meta_description": "Base para guarda-sol e ombrellone com 18kg preenchida com água ou areia — suporte estável para uso na praia. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/cadeira-espreguicadeira-em-aluminio-marrom-414712-bel-77702",
        "keyword": "espreguiçadeira alumínio",
        "seo_title": "Cadeira Espreguiçadeira Alumínio Marrom Bel | Ultra Máquinas",
        "meta_description": "Espreguiçadeira em alumínio marrom resistente — ideal para piscina, jardim e área de lazer com design leve. Compre no Pix e ganhe 10% de desconto.",
    },
    {
        "url": "https://www.ultramaquinas.com.br/produto/caixa-termica-branca-113-litros-igloo-279411-kala-81054",
        "keyword": "caixa térmica igloo 113 litros",
        "seo_title": "Caixa Térmica Igloo 113 Litros Branca Kala | Ultra Máquinas",
        "meta_description": "Caixa térmica Igloo 113L — alto isolamento, dreno de fácil acesso e tampa robusta para pescaria e acampamento. Compre no Pix e ganhe 10% de desconto.",
    },
]


def export_excel(products, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SEO Metadata"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "URL", "Keyword Principal", "SEO Title", "Chars Title", "Meta Description", "Chars Meta"]
    col_widths = [4, 60, 30, 62, 12, 165, 12]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w / 7

    alt_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    wrap = Alignment(vertical="top", wrap_text=True)
    center = Alignment(horizontal="center", vertical="top")

    for i, p in enumerate(products, 1):
        row = i + 1
        fill = alt_fill if i % 2 == 0 else white_fill

        title_len = len(p["seo_title"])
        meta_len = len(p["meta_description"])

        title_warn_font = Font(color="C00000", bold=True) if title_len > 60 else Font()
        meta_warn_font = Font(color="C00000", bold=True) if meta_len > 160 else Font()

        data = [i, p["url"], p["keyword"], p["seo_title"], title_len, p["meta_description"], meta_len]
        aligns = [center, wrap, wrap, wrap, center, wrap, center]

        for col, (val, align) in enumerate(zip(data, aligns), 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill
            cell.alignment = align
            cell.border = border
            if col == 5:
                cell.font = title_warn_font
            elif col == 7:
                cell.font = meta_warn_font

    ws.row_dimensions[1].height = 30
    for i in range(2, len(products) + 2):
        ws.row_dimensions[i].height = 45

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{len(products) + 1}"

    wb.save(path)
    print(f"Excel salvo em: {path}")
    return path


if __name__ == "__main__":
    import os
    os.makedirs("output", exist_ok=True)
    export_excel(PRODUCTS, "output/seo_metadata.xlsx")

    print("\nResumo:")
    for p in PRODUCTS:
        t = len(p["seo_title"])
        m = len(p["meta_description"])
        t_ok = "✓" if t <= 60 else f"⚠ {t}"
        m_ok = "✓" if 140 <= m <= 160 else f"⚠ {m}"
        print(f"  [{t_ok}] [{m_ok}] {p['seo_title'][:55]}")
